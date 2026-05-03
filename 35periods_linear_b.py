"""
Assignment 6b: Realized demand evaluation (35 periods, linear extrapolation)

"""

import math
import gurobipy as gp
from gurobipy import GRB
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
from input_data_35periods_linear import (
    PARTS, END_PRODUCT, T, BOM, LEAD_TIME, MIN_LOT,
    INIT_INV, SETUP_COST, HOLDING_COST,
    DEMAND_FORECAST, DEMAND_REALIZED, BACKORDER_COST
)

def get_parents(part):
    parents = {}
    for parent, children in BOM.items():
        if part in children:
            parents[parent] = children[part]
    return parents

def clean_num(val, tol=1e-6):
    v = float(val)
    return 0.0 if abs(v) < tol else round(v, 6)

#  Rescaling factor: 30 original weeks / 35 extended weeks 
D = 30 / 35
EVAL_WEEKS = 30   # service metrics computed on the first 30 weeks only

#  Capacity / cost parameters
CAP_X_BASE     = 800
CAP_X_MAX_EXP  = 200
COST_EXP_X     = 10
CAP_X_OT_MAX   = 300
COST_OT_X      = 2
CAP_Y_BASE     = 60 * 24 * 7 - 80
CAP_Y_MAX_PCT  = 40
COST_EXP_Y_PCT = 1_500
CAP_Y_OT_MAX   = 38
COST_OT_Y      = 120
PROC_X = {END_PRODUCT: 1}
PROC_Y = {"B1401": 3, "B2302": 2}

periods = range(1, T + 1)

# Step 1: Re-solve 6a to get the fixed 35-week plan
m6a = gp.Model("APM_6a_reload")
m6a.setParam("OutputFlag", 0)
xa     = m6a.addVars(PARTS, periods, name="x",  vtype=GRB.INTEGER, lb=0)
ya     = m6a.addVars(PARTS, periods, name="y",  vtype=GRB.BINARY)
Ia     = m6a.addVars(PARTS, periods, name="I",  vtype=GRB.INTEGER, lb=0)
dxa    = m6a.addVar(name="dx",     vtype=GRB.INTEGER, lb=0, ub=CAP_X_MAX_EXP)
dya    = m6a.addVar(name="dy_pct", lb=0, ub=CAP_Y_MAX_PCT)
ot_xa  = m6a.addVars(periods, name="ot_x", vtype=GRB.INTEGER, lb=0, ub=CAP_X_OT_MAX)
ot_ya  = m6a.addVars(periods, name="ot_y", lb=0, ub=CAP_Y_OT_MAX)
BIG_M  = {i: sum(DEMAND_FORECAST)*25 for i in PARTS}

m6a.setObjective(
    gp.quicksum(SETUP_COST[i]*ya[i,t] + HOLDING_COST[i]*Ia[i,t]
                for i in PARTS for t in periods)
    + COST_EXP_X*dxa + COST_EXP_Y_PCT*dya
    + gp.quicksum(COST_OT_X*ot_xa[t] + COST_OT_Y*ot_ya[t] for t in periods),
    GRB.MINIMIZE)

for i in PARTS:
    for t in periods:
        inv_prev = INIT_INV[i] if t==1 else Ia[i,t-1]
        op = t - LEAD_TIME[i]
        rec = xa[i,op] if op>=1 else 0
        int_d = gp.quicksum(BOM[p][i]*xa[p,t] for p in get_parents(i))
        ext_d = DEMAND_FORECAST[t-1] if i==END_PRODUCT else 0
        m6a.addConstr(Ia[i,t] == inv_prev + rec - int_d - ext_d)
        m6a.addConstr(xa[i,t] >= MIN_LOT[i]*ya[i,t])
        m6a.addConstr(xa[i,t] <= BIG_M[i]*ya[i,t])

for t in periods:
    m6a.addConstr(gp.quicksum(PROC_X.get(i,0)*xa[i,t] for i in PARTS)
                  <= CAP_X_BASE + dxa + ot_xa[t])
    m6a.addConstr(gp.quicksum(PROC_Y.get(i,0)*xa[i,t] for i in PARTS)
                  <= CAP_Y_BASE + (CAP_Y_BASE/100)*dya + 60*ot_ya[t])
m6a.optimize()

dx_val     = clean_num(dxa.X)
dy_pct_val = clean_num(dya.X)
invest_X   = COST_EXP_X     * dx_val
invest_Y   = COST_EXP_Y_PCT * dy_pct_val
cap_x_new  = CAP_X_BASE + dx_val
cap_y_new  = CAP_Y_BASE * (1 + dy_pct_val/100)
cost_6a    = m6a.ObjVal
total_ot_x_6a = sum(COST_OT_X * ot_xa[t].X for t in periods)
total_ot_y_6a = sum(COST_OT_Y * ot_ya[t].X for t in periods)
ot_x_vals  = {t: clean_num(ot_xa[t].X) for t in periods}
ot_y_vals  = {t: clean_num(ot_ya[t].X) for t in periods}

schedule = {i: {t: xa[i,t].X for t in periods} for i in PARTS}
setup_cost_fixed = sum(SETUP_COST[i] for i in PARTS for t in periods if xa[i,t].X > 0.5)

# Step 2: Simulate with realized demand (35 weeks)
inventory = {i: {} for i in PARTS}
held      = {i: {} for i in PARTS}
backorders = {t: 0.0 for t in periods}
net_end    = {}

for i in PARTS:
    if i == END_PRODUCT:
        continue
    inv_prev = float(INIT_INV[i])
    for t in periods:
        op = t - LEAD_TIME[i]
        rec = schedule[i].get(op, 0.0) if op >= 1 else 0.0
        int_d = sum(BOM[p][i]*schedule[p].get(t,0.0) for p in get_parents(i))
        net = inv_prev + rec - int_d
        inventory[i][t] = clean_num(net)
        held[i][t]      = clean_num(max(0.0, net))
        inv_prev = net

inv_prev = float(INIT_INV[END_PRODUCT])
bo_prev  = 0.0
for t in periods:
    op = t - LEAD_TIME[END_PRODUCT]
    rec = schedule[END_PRODUCT].get(op, 0.0) if op >= 1 else 0.0
    net = inv_prev - bo_prev + rec - DEMAND_REALIZED[t-1]
    net_end[t] = clean_num(net)
    if net >= 0:
        inventory[END_PRODUCT][t] = clean_num(net)
        held[END_PRODUCT][t]      = clean_num(net)
        backorders[t] = 0.0
    else:
        inventory[END_PRODUCT][t] = 0.0
        held[END_PRODUCT][t]      = 0.0
        backorders[t] = clean_num(-net)
    inv_prev = inventory[END_PRODUCT][t]
    bo_prev  = backorders[t]

# Step 3: Compute costs (35-week raw, then 30-week equivalent)
total_holding_35   = sum(HOLDING_COST[i]*held[i][t] for i in PARTS for t in periods)
total_backorder_35 = sum(BACKORDER_COST * backorders[t] for t in periods)
total_cost_6b_35   = (setup_cost_fixed + total_holding_35 + total_backorder_35
                     + invest_X + invest_Y + total_ot_x_6a + total_ot_y_6a)

# Rescale to 30-week equivalent (investments NOT scaled, backorder ceil)
setup_30      = setup_cost_fixed * D
holding_30    = total_holding_35 * D
backorder_30  = math.ceil(total_backorder_35 * D)   
ot_x_30       = total_ot_x_6a * D
ot_y_30       = total_ot_y_6a * D
total_cost_6b_30 = (setup_30 + holding_30 + backorder_30
                   + invest_X + invest_Y + ot_x_30 + ot_y_30)

# Step 4: Service metrics on FIRST 30 WEEKS (for fair comparison)
eval_periods = range(1, EVAL_WEEKS + 1)

periods_no_bo = sum(1 for t in eval_periods if backorders[t] == 0)
service_level = periods_no_bo / EVAL_WEEKS

total_demand_eval = sum(DEMAND_REALIZED[t-1] for t in eval_periods)
total_new_bo = 0.0
prev_bo = 0.0
for t in eval_periods:
    new_bo = max(0.0, backorders[t] - prev_bo)
    total_new_bo += new_bo
    prev_bo = backorders[t]
fill_rate     = 1.0 - (total_new_bo / total_demand_eval) if total_demand_eval > 0 else 1.0
units_on_time = total_demand_eval - total_new_bo
total_new_bo_ceil = math.ceil(total_new_bo)   # conservative rounding

# Step 5: Read 5b reference for comparison
def read_float(ws, row, col=3):
    try: return float(ws.cell(row, col).value)
    except: return None

import re
def parse_pct(s):
    m = re.search(r"[\d.]+", str(s) if s else "")
    return float(m.group()) if m else None

ref_5b = None
try:
    wb_ref = load_workbook("OUTPUT.xlsx", data_only=True)
    if "Output_5b" in wb_ref.sheetnames:
        ws5b = wb_ref["Output_5b"]
        ref_5b = dict(
            setup    = read_float(ws5b, 3),
            holding  = read_float(ws5b, 4),
            bo       = read_float(ws5b, 5),
            invest_X = read_float(ws5b, 6),
            invest_Y = read_float(ws5b, 7),
            ot_x     = read_float(ws5b, 8),
            ot_y     = read_float(ws5b, 9),
            total    = read_float(ws5b, 10),
            sl       = parse_pct(ws5b.cell(15, 3).value),
            fr       = parse_pct(ws5b.cell(16, 3).value),
        )
except Exception:
    ref_5b = None

# Step 6: Write to OUTPUT.xlsx
OUTPUT_FILE = "OUTPUT.xlsx"
SHEET_NAME  = "Output_linear_realdemand"

if os.path.exists(OUTPUT_FILE):
    wb = load_workbook(OUTPUT_FILE)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)
else:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

NO_FILL    = PatternFill(fill_type=None)
BLACK_FILL = PatternFill("solid", start_color="000000", end_color="000000")
none_border = Border()
bot_medium  = Border(bottom=Side(style="medium", color="000000"))
bot_thin    = Border(bottom=Side(style="thin",   color="CCCCCC"))

def plain(cell, val, bold=False, align="left", size=9, color="000000", fmt=None, border=None):
    cell.value = val
    cell.font  = Font(name="Calibri", bold=bold, size=size, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.fill  = NO_FILL
    if border: cell.border = border
    if fmt:    cell.number_format = fmt

def hdr(cell, val):
    cell.value = val
    cell.font  = Font(name="Calibri", bold=False, size=9, color="FFFFFF")
    cell.fill  = BLACK_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = none_border

def section_title(ws, r, last_col, text):
    ws.row_dimensions[r].height = 14
    plain(ws.cell(r, 2), text, size=8, color="888888", border=bot_medium)
    for col in range(3, last_col + 1):
        ws.cell(r, col).border = bot_medium

ws.column_dimensions["A"].width = 1
ws.column_dimensions["B"].width = 9
for col in range(3, T + 4):
    ws.column_dimensions[get_column_letter(col)].width = 5.5
last_col = T + 2

#  Title 
ws.row_dimensions[1].height = 26
ws.merge_cells(f"B1:{get_column_letter(last_col)}1")
plain(ws.cell(1, 2),
          "Assignment 6a - Extended horizon (35 periods, linear extrapolation) - Realized demand evaluation",
      bold=True, size=13)

#  Cost summary (35-week raw) 
ws.row_dimensions[2].height = 4
cost_rows = [
    ("Setup cost (35w)",      setup_cost_fixed,    False),
    ("Holding cost (35w)",    total_holding_35,    False),
    ("Backorder cost (35w)",  total_backorder_35,  False),
    ("Investment cost X",     invest_X,            False),
    ("Investment cost Y",     invest_Y,            False),
    ("Overtime cost X (35w)", total_ot_x_6a,       False),
    ("Overtime cost Y (35w)", total_ot_y_6a,       False),
    ("Total cost (35w)",      total_cost_6b_35,    True),
]
for r, (label, val, bold) in enumerate(cost_rows, start=3):
    ws.row_dimensions[r].height = 16
    plain(ws.cell(r, 2), label, size=9, color="555555")
    plain(ws.cell(r, 3), round(clean_num(val), 2), bold=bold, size=9, fmt='"€"#,##0.00')
    ws.merge_cells(f"C{r}:{get_column_letter(last_col)}{r}")

#  Rescaled costs (30-week equivalent) 
rescaled_rows = [
    ("Setup cost (30w eq.)",       setup_30,         False),
    ("Holding cost (30w eq.)",     holding_30,       False),
    ("Backorder cost (30w eq.)",   backorder_30,     False),
    ("Overtime cost X (30w eq.)",  ot_x_30,          False),
    ("Overtime cost Y (30w eq.)",  ot_y_30,          False),
    ("Investment X (not scaled)",  invest_X,         False),
    ("Investment Y (not scaled)",  invest_Y,         False),
    ("Total cost (30w eq.)",       total_cost_6b_30, True),
]
start_row = 3 + len(cost_rows) + 1
plain(ws.cell(start_row, 2),
      "Rescaled to 30-week equivalent  (factor D = 30/35; backorders rounded up)",
      size=8, color="888888")
start_row += 1
for r, (label, val, bold) in enumerate(rescaled_rows, start=start_row):
    ws.row_dimensions[r].height = 16
    plain(ws.cell(r, 2), label, size=9, color="555555")
    plain(ws.cell(r, 3), round(clean_num(val), 2), bold=bold, size=9, fmt='"€"#,##0.00')
    ws.merge_cells(f"C{r}:{get_column_letter(last_col)}{r}")

#  Service metrics 
metrics_start = start_row + len(rescaled_rows) + 1
section_title(ws, metrics_start, last_col,
              f"Service metrics (end product, evaluated on first {EVAL_WEEKS} weeks)")
metric_rows = [
    ("Service level",   f"{service_level*100:.1f}%  ({periods_no_bo}/{EVAL_WEEKS} periods without backorder)", False),
    ("Fill rate",       f"{fill_rate*100:.2f}%  ({units_on_time:,.0f} / {total_demand_eval:,.0f} units on time)", False),
    ("Total backorder", f"{total_new_bo_ceil:,.0f} units (ceil of {total_new_bo:.1f})", total_new_bo > 0),
]
for r, (label, text, red) in enumerate(metric_rows, start=metrics_start + 1):
    ws.row_dimensions[r].height = 16
    plain(ws.cell(r, 2), label, size=9, color="555555")
    plain(ws.cell(r, 3), text, size=9, color="CC0000" if red else "000000")
    ws.merge_cells(f"C{r}:{get_column_letter(last_col)}{r}")

# Comparison table 5b vs 6b (30w equivalent)
r = metrics_start + len(metric_rows) + 2
section_title(ws, r, last_col, "Comparison: 5b vs 6b (30-week equivalent)")
r += 1
ws.row_dimensions[r].height = 16

for c, w in [(2, 22), (3, 14), (4, 14), (5, 14), (6, 10)]:
    ws.column_dimensions[get_column_letter(c)].width = w

hdr(ws.cell(r, 2), "Metric")
hdr(ws.cell(r, 3), "5b")
hdr(ws.cell(r, 4), "6b (30w eq.)")
hdr(ws.cell(r, 5), "Diff")
hdr(ws.cell(r, 6), "Diff %")

vals_6b = dict(
    setup    = setup_30,
    holding  = holding_30,
    bo       = backorder_30,
    invest_X = invest_X,
    invest_Y = invest_Y,
    ot_x     = ot_x_30,
    ot_y     = ot_y_30,
    total    = total_cost_6b_30,
    sl       = service_level * 100,
    fr       = fill_rate * 100,
)

cmp_rows = [
    ("Setup cost (EUR)",     "setup",    '"€"#,##0.00', False),
    ("Holding cost (EUR)",   "holding",  '"€"#,##0.00', False),
    ("Backorder cost (EUR)", "bo",       '"€"#,##0.00', False),
    ("Investment X (EUR)",   "invest_X", '"€"#,##0.00', False),
    ("Investment Y (EUR)",   "invest_Y", '"€"#,##0.00', False),
    ("Overtime X (EUR)",     "ot_x",     '"€"#,##0.00', False),
    ("Overtime Y (EUR)",     "ot_y",     '"€"#,##0.00', False),
    ("Total cost (EUR)",     "total",    '"€"#,##0.00', True),
    ("Service level (%)",    "sl",       '0.0',         False),
    ("Fill rate (%)",        "fr",       '0.00',        False),
]

for label, key, fmt, is_total in cmp_rows:
    r += 1
    ws.row_dimensions[r].height = 15
    v6 = vals_6b[key]
    v5 = ref_5b.get(key) if ref_5b else None
    plain(ws.cell(r, 2), label, size=9, bold=is_total, border=bot_thin)
    if v5 is not None:
        plain(ws.cell(r, 3), round(float(v5), 2), size=9, bold=is_total,
              align="right", fmt=fmt, border=bot_thin)
        plain(ws.cell(r, 4), round(v6, 2), size=9, bold=is_total,
              align="right", fmt=fmt, border=bot_thin)
        diff = v6 - float(v5)
        pct = (diff / float(v5)) if float(v5) != 0 else 0
        
        is_kpi = key in ("sl", "fr")
        good = (diff < 0) if not is_kpi else (diff > 0)
        bad  = (diff > 0) if not is_kpi else (diff < 0)
        color = "117733" if good else ("CC0000" if bad else "000000")
        diff_fmt = '"€"#,##0.00;[Red]-"€"#,##0.00' if "EUR" in label else '+0.00;-0.00;0.00'
        plain(ws.cell(r, 5), round(diff, 2), size=9, bold=is_total,
              align="right", color=color, fmt=diff_fmt, border=bot_thin)
        plain(ws.cell(r, 6), pct, size=9, bold=is_total,
              align="right", color=color, fmt='+0.0%;-0.0%;0.0%', border=bot_thin)
    else:
        plain(ws.cell(r, 3), "n/a", size=9, align="right", color="AAAAAA", border=bot_thin)
        plain(ws.cell(r, 4), round(v6, 2), size=9, bold=is_total,
              align="right", fmt=fmt, border=bot_thin)
        plain(ws.cell(r, 5), "", size=9, border=bot_thin)
        plain(ws.cell(r, 6), "", size=9, border=bot_thin)

if ref_5b is None:
    r += 1
    plain(ws.cell(r, 2),
          "(5b reference not found - run assignment_5b.py first to enable comparison.)",
          size=9, color="AA0000")

# Reset week column widths
ws.column_dimensions["B"].width = 9
for col in range(3, T + 4):
    ws.column_dimensions[get_column_letter(col)].width = 5.5

# Backorder schedule
r += 2
section_title(ws, r, last_col, f"Backorder schedule - {END_PRODUCT} (end product only)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "Part")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))

r += 1
ws.row_dimensions[r].height = 15
plain(ws.cell(r, 2), END_PRODUCT, size=9, border=bot_thin)
for t in periods:
    v = round(backorders[t])
    if v > 0:
        plain(ws.cell(r, t+2), v, bold=True, size=9, align="center",
              color="CC0000", fmt='#,##0', border=bot_thin)
    else:
        plain(ws.cell(r, t+2), "", size=9, align="center", border=bot_thin)

# Demand vs delivered
r += 2
section_title(ws, r, last_col, "Demand vs delivered (end product)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))

r += 1
ws.row_dimensions[r].height = 15
plain(ws.cell(r, 2), "Demand", size=9, border=bot_thin)
for t in periods:
    plain(ws.cell(r, t+2), DEMAND_REALIZED[t-1], size=9, align="center",
          fmt='#,##0', border=bot_thin)

r += 1
ws.row_dimensions[r].height = 15
plain(ws.cell(r, 2), "Delivered", size=9, border=bot_thin)
prev_bo = 0.0
for t in periods:
    new_bo_t = max(0.0, backorders[t] - prev_bo)
    delivered = DEMAND_REALIZED[t-1] - new_bo_t
    shortage  = delivered < DEMAND_REALIZED[t-1]
    plain(ws.cell(r, t+2), round(delivered), size=9, align="center",
          color="CC0000" if shortage else "000000",
          fmt='#,##0', border=bot_thin)
    prev_bo = backorders[t]

# Production schedule
r += 2
section_title(ws, r, last_col, "Production / order schedule (units, from 6a plan)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "Part")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))

for i in PARTS:
    r += 1
    ws.row_dimensions[r].height = 15
    plain(ws.cell(r, 2), i, size=9, border=bot_thin)
    for t in periods:
        val = round(schedule[i][t]) if schedule[i][t] > 0.5 else ""
        plain(ws.cell(r, t+2), val, bold=bool(val), size=9, align="center",
              fmt='#,##0' if val != "" else None, border=bot_thin)

# Inventory levels (realized)
r += 2
section_title(ws, r, last_col, "Inventory levels (end of period, realized)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "Part")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))

for i in PARTS:
    r += 1
    ws.row_dimensions[r].height = 15
    plain(ws.cell(r, 2), i, size=9, border=bot_thin)
    for t in periods:
        v = round(inventory[i][t])
        plain(ws.cell(r, t+2), v, size=9, align="center",
              color="BBBBBB" if v == 0 else "000000",
              fmt='#,##0', border=bot_thin)

# Overtime usage (from 6a plan)
r += 2
section_title(ws, r, last_col, "Overtime usage (from 6a plan, unchanged)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))

for label, vals, fmt in [
    ("X  (units OT)", ot_x_vals, '#,##0'),
    ("Y  (hours OT)", ot_y_vals, '0.00'),
]:
    r += 1
    ws.row_dimensions[r].height = 15
    plain(ws.cell(r, 2), label, size=9, border=bot_thin)
    for t in periods:
        v = round(vals[t], 2)
        col = "CC0000" if v else "BBBBBB"
        plain(ws.cell(r, t+2), v if v else "", bold=bool(v), size=9,
              align="center", color=col,
              fmt=fmt if v else None, border=bot_thin)

# Setup decisions
r += 2
section_title(ws, r, last_col, "Setup decisions")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "Part")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))

for i in PARTS:
    r += 1
    ws.row_dimensions[r].height = 15
    plain(ws.cell(r, 2), i, size=9, border=bot_thin)
    for t in periods:
        plain(ws.cell(r, t+2), "x" if schedule[i][t] > 0.5 else "",
              bold=True, size=9, align="center", border=bot_thin)

wb.save(OUTPUT_FILE)
print(f"Results written to {OUTPUT_FILE} -> sheet '{SHEET_NAME}'")
print()
print(f"=== 35-week horizon (raw) ===")
print(f"Total cost (6b): EUR {clean_num(total_cost_6b_35):,.2f}")
print(f"  Setup:         EUR {clean_num(setup_cost_fixed):,.2f}")
print(f"  Holding:       EUR {clean_num(total_holding_35):,.2f}")
print(f"  Backorder:     EUR {clean_num(total_backorder_35):,.2f}")
print(f"  Investment X:  EUR {clean_num(invest_X):,.2f}")
print(f"  Investment Y:  EUR {clean_num(invest_Y):,.2f}")
print(f"  Overtime X:    EUR {clean_num(total_ot_x_6a):,.2f}")
print(f"  Overtime Y:    EUR {clean_num(total_ot_y_6a):,.2f}")
print()
print(f"=== 30-week equivalent (D = 30/35; investments NOT scaled, BO ceiled) ===")
print(f"Total cost:      EUR {total_cost_6b_30:,.2f}")
print(f"  Setup:         EUR {setup_30:,.2f}")
print(f"  Holding:       EUR {holding_30:,.2f}")
print(f"  Backorder:     EUR {backorder_30:,.2f}  (ceil of {total_backorder_35*D:.2f})")
print(f"  Investment X:  EUR {invest_X:,.2f}  (not scaled)")
print(f"  Investment Y:  EUR {invest_Y:,.2f}  (not scaled)")
print(f"  Overtime X:    EUR {ot_x_30:,.2f}")
print(f"  Overtime Y:    EUR {ot_y_30:,.2f}")
print()
print(f"Service level:   {service_level*100:.1f}%  ({periods_no_bo}/{EVAL_WEEKS} periods, first {EVAL_WEEKS} weeks)")
print(f"Fill rate:       {fill_rate*100:.2f}%  ({units_on_time:,.0f}/{total_demand_eval:,.0f} units, first {EVAL_WEEKS} weeks)")
print(f"Backorder units: {total_new_bo_ceil} (ceil of {total_new_bo:.1f})")

if ref_5b is not None:
    print()
    print(f"=== Comparison 5b vs 6b (30-week equivalent) ===")
    print(f"{'Metric':<22}{'5b':>14}{'6b':>14}{'Diff':>14}{'Diff %':>10}")
    print("-" * 74)
    comp = [
        ("Setup",        ref_5b["setup"],    setup_30),
        ("Holding",      ref_5b["holding"],  holding_30),
        ("Backorder",    ref_5b["bo"],       backorder_30),
        ("Invest X",     ref_5b["invest_X"], invest_X),
        ("Invest Y",     ref_5b["invest_Y"], invest_Y),
        ("Overtime X",   ref_5b["ot_x"],     ot_x_30),
        ("Overtime Y",   ref_5b["ot_y"],     ot_y_30),
        ("Total cost",   ref_5b["total"],    total_cost_6b_30),
    ]
    for label, v5, v6 in comp:
        if v5 is None: continue
        diff = v6 - float(v5)
        pct = (diff / float(v5) * 100) if float(v5) != 0 else 0
        print(f"{label:<22}{float(v5):>14,.2f}{v6:>14,.2f}{diff:>+14,.2f}{pct:>+9.1f}%")
    if ref_5b.get("sl") is not None and ref_5b.get("fr") is not None:
        print(f"{'Service level (%)':<22}{ref_5b['sl']:>14.1f}{service_level*100:>14.1f}{(service_level*100-ref_5b['sl']):>+14.1f}")
        print(f"{'Fill rate (%)':<22}{ref_5b['fr']:>14.2f}{fill_rate*100:>14.2f}{(fill_rate*100-ref_5b['fr']):>+14.2f}")
else:
    print()
    print("(Run assignment_5b.py first to enable side-by-side comparison.)")