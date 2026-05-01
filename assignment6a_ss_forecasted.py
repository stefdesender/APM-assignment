"""
APM Project 2026 - Assignment 6a (forecasted demand, with safety stock)

Runs two scenarios:
  - 95% cycle service level (z = 1.65)
  - 99% cycle service level (z = 2.33)

For each scenario:
  1. Compute safety stock from forecast errors (sigma_e, L+R, z).
  2. Solve MIP on FORECASTED demand with I[E2801,t] >= SS.
  3. Simulate plan against FORECASTED demand (verification).
  4. Write results to separate Excel sheet + print to terminal.

At the end, a side-by-side comparison of both scenarios is printed.
"""

import math
import statistics
import gurobipy as gp
from gurobipy import GRB
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
from input_data import (
    PARTS, END_PRODUCT, T, BOM, LEAD_TIME, MIN_LOT,
    INIT_INV, SETUP_COST, HOLDING_COST,
    DEMAND_FORECAST, DEMAND_REALIZED, BACKORDER_COST
)

# ── Helper functions ───────────────────────────────────────────────────────
def get_parents(part):
    parents = {}
    for parent, children in BOM.items():
        if part in children:
            parents[parent] = children[part]
    return parents

def clean_num(val, tol=1e-6):
    v = float(val)
    return 0.0 if abs(v) < tol else round(v, 6)

# ── Shared safety stock parameters ────────────────────────────────────────
forecast_errors = [DEMAND_REALIZED[t] - DEMAND_FORECAST[t] for t in range(T)]
sigma_e   = statistics.stdev(forecast_errors)
L         = LEAD_TIME[END_PRODUCT]
R         = 1
sigma_LTD = sigma_e * math.sqrt(L + R)

print(f"\n── Safety Stock Base Parameters ──────────────────────")
print(f"  sigma_e (per period):     {sigma_e:.2f} units")
print(f"  Lead time L:              {L} period(s)")
print(f"  Review period R:          {R} period(s)")
print(f"  sigma_LTD (L+R={L+R}):     {sigma_LTD:.2f} units")
print(f"──────────────────────────────────────────────────────\n")

# ── Capacity parameters ──────────────────────────────────────────────────
CAP_X_BASE     = 800
CAP_X_MAX_EXP  = 200
COST_EXP_X     = 10
CAP_X_OT_MAX   = 300
COST_OT_X      = 2
CAP_Y_BASE     = 60 * 24 * 7 - 80
CAP_Y_MAX_PCT  = 40
COST_EXP_Y_PCT = 1_500
CAP_Y_OT_MAX   = 38
COST_OT_Y      = 120
PROC_X = {END_PRODUCT: 1}
PROC_Y = {"B1401": 3, "B2302": 2}

periods = range(1, T + 1)

# ── Scenarios to run ─────────────────────────────────────────────────────
SCENARIOS = [
    {"label": "95% CSL", "z": 1.65, "sheet": "Output_6a_ss_fc_95"},
    {"label": "99% CSL", "z": 2.33, "sheet": "Output_6a_ss_fc_99"},
]

OUTPUT_FILE = "OUTPUT.xlsx"
results = []

for scenario in SCENARIOS:
    z_val  = scenario["z"]
    label  = scenario["label"]
    sheet  = scenario["sheet"]

    SS_raw       = z_val * sigma_LTD
    SAFETY_STOCK = math.ceil(SS_raw)

    print(f"{'='*60}")
    print(f"  SCENARIO: {label}  (z = {z_val},  SS = {SAFETY_STOCK} units)")
    print(f"{'='*60}")

    # ── Build and solve MIP on FORECASTED demand ──────────────────────
    m = gp.Model(f"APM_6a_ss_{label}")
    m.setParam("OutputFlag", 0)

    x    = m.addVars(PARTS, periods, name="x",    vtype=GRB.INTEGER, lb=0)
    y    = m.addVars(PARTS, periods, name="y",    vtype=GRB.BINARY)
    I    = m.addVars(PARTS, periods, name="I",    vtype=GRB.INTEGER, lb=0)
    dx   = m.addVar(name="dx",                    vtype=GRB.INTEGER, lb=0, ub=CAP_X_MAX_EXP)
    dy   = m.addVar(name="dy_pct",                lb=0,              ub=CAP_Y_MAX_PCT)
    ot_x = m.addVars(periods, name="ot_x",        vtype=GRB.INTEGER, lb=0, ub=CAP_X_OT_MAX)
    ot_y = m.addVars(periods, name="ot_y",        lb=0,              ub=CAP_Y_OT_MAX)

    BIG_M = {i: sum(DEMAND_FORECAST) * 25 for i in PARTS}

    m.setObjective(
        gp.quicksum(SETUP_COST[i] * y[i, t] + HOLDING_COST[i] * I[i, t]
                    for i in PARTS for t in periods)
        + COST_EXP_X * dx
        + COST_EXP_Y_PCT * dy
        + gp.quicksum(COST_OT_X * ot_x[t] + COST_OT_Y * ot_y[t] for t in periods),
        GRB.MINIMIZE
    )

    for i in PARTS:
        for t in periods:
            inv_prev = INIT_INV[i] if t == 1 else I[i, t - 1]
            op  = t - LEAD_TIME[i]
            rec = x[i, op] if op >= 1 else 0
            int_d = gp.quicksum(BOM[p][i] * x[p, t] for p in get_parents(i))
            ext_d = DEMAND_FORECAST[t - 1] if i == END_PRODUCT else 0
            m.addConstr(I[i, t] == inv_prev + rec - int_d - ext_d)
            m.addConstr(x[i, t] >= MIN_LOT[i] * y[i, t])
            m.addConstr(x[i, t] <= BIG_M[i]   * y[i, t])

    # Safety stock constraint
    for t in periods:
        m.addConstr(I[END_PRODUCT, t] >= SAFETY_STOCK)

    # Capacity constraints
    for t in periods:
        m.addConstr(
            gp.quicksum(PROC_X.get(i, 0) * x[i, t] for i in PARTS)
            <= CAP_X_BASE + dx + ot_x[t])
        m.addConstr(
            gp.quicksum(PROC_Y.get(i, 0) * x[i, t] for i in PARTS)
            <= CAP_Y_BASE + (CAP_Y_BASE / 100) * dy + 60 * ot_y[t])

    m.optimize()

    if m.Status != GRB.OPTIMAL:
        print(f"  Model status: {m.Status}. No optimal solution found.")
        continue

    # ── Extract plan ──────────────────────────────────────────────────
    dx_val      = clean_num(dx.X)
    dy_pct_val  = clean_num(dy.X)
    invest_X    = COST_EXP_X     * dx_val
    invest_Y    = COST_EXP_Y_PCT * dy_pct_val
    cap_x_new   = CAP_X_BASE + dx_val
    cap_y_new   = CAP_Y_BASE * (1 + dy_pct_val / 100)
    total_ot_x  = sum(COST_OT_X * ot_x[t].X for t in periods)
    total_ot_y  = sum(COST_OT_Y * ot_y[t].X for t in periods)
    ot_x_vals   = {t: clean_num(ot_x[t].X) for t in periods}
    ot_y_vals   = {t: clean_num(ot_y[t].X) for t in periods}

    schedule        = {i: {t: x[i, t].X for t in periods} for i in PARTS}
    setup_cost_plan = sum(SETUP_COST[i] for i in PARTS for t in periods if x[i, t].X > 0.5)

    # ── Simulate against FORECASTED demand ────────────────────────────
    inventory  = {i: {} for i in PARTS}
    held       = {i: {} for i in PARTS}
    backorders = {t: 0.0 for t in periods}
    net_end    = {}

    for i in PARTS:
        if i == END_PRODUCT:
            continue
        inv_prev = float(INIT_INV[i])
        for t in periods:
            op    = t - LEAD_TIME[i]
            rec   = schedule[i].get(op, 0.0) if op >= 1 else 0.0
            int_d = sum(BOM[p][i] * schedule[p].get(t, 0.0) for p in get_parents(i))
            net   = inv_prev + rec - int_d
            inventory[i][t] = clean_num(net)
            held[i][t]      = clean_num(max(0.0, net))
            inv_prev = net

    inv_prev = float(INIT_INV[END_PRODUCT])
    bo_prev  = 0.0
    for t in periods:
        op  = t - LEAD_TIME[END_PRODUCT]
        rec = schedule[END_PRODUCT].get(op, 0.0) if op >= 1 else 0.0
        net = inv_prev - bo_prev + rec - DEMAND_FORECAST[t - 1]
        net_end[t] = clean_num(net)
        if net >= 0:
            inventory[END_PRODUCT][t] = clean_num(net)
            held[END_PRODUCT][t]      = clean_num(net)
            backorders[t] = 0.0
        else:
            inventory[END_PRODUCT][t] = 0.0
            held[END_PRODUCT][t]      = 0.0
            backorders[t] = clean_num(-net)
        inv_prev = inventory[END_PRODUCT][t]
        bo_prev  = backorders[t]

    # ── Costs ─────────────────────────────────────────────────────────
    total_holding   = sum(HOLDING_COST[i] * held[i][t] for i in PARTS for t in periods)
    total_backorder = sum(BACKORDER_COST * backorders[t] for t in periods)
    total_cost      = (setup_cost_plan + total_holding + total_backorder
                       + invest_X + invest_Y + total_ot_x + total_ot_y)

    # ── Service metrics ───────────────────────────────────────────────
    periods_no_bo = sum(1 for t in periods if backorders[t] == 0)
    service_level = periods_no_bo / T
    total_demand  = sum(DEMAND_FORECAST)
    total_new_bo  = 0.0
    prev_bo = 0.0
    for t in periods:
        new_bo = max(0.0, backorders[t] - prev_bo)
        total_new_bo += new_bo
        prev_bo = backorders[t]
    fill_rate     = 1.0 - (total_new_bo / total_demand) if total_demand > 0 else 1.0
    units_on_time = total_demand - total_new_bo

    # ── Store for comparison ──────────────────────────────────────────
    results.append({
        "label":       label,
        "z":           z_val,
        "ss":          SAFETY_STOCK,
        "setup":       setup_cost_plan,
        "holding":     total_holding,
        "backorder":   total_backorder,
        "invest_X":    invest_X,
        "invest_Y":    invest_Y,
        "ot_x":        total_ot_x,
        "ot_y":        total_ot_y,
        "total":       total_cost,
        "sl":          service_level,
        "fr":          fill_rate,
        "bo_units":    total_new_bo,
    })

    # ── Console output ────────────────────────────────────────────────
    print(f"\n  Cost Summary (forecasted demand):")
    print(f"    Total cost:      EUR {clean_num(total_cost):>12,.2f}")
    print(f"      Setup:         EUR {clean_num(setup_cost_plan):>12,.2f}")
    print(f"      Holding:       EUR {clean_num(total_holding):>12,.2f}")
    print(f"      Backorder:     EUR {clean_num(total_backorder):>12,.2f}")
    print(f"      Investment X:  EUR {clean_num(invest_X):>12,.2f}")
    print(f"      Investment Y:  EUR {clean_num(invest_Y):>12,.2f}")
    print(f"      Overtime X:    EUR {clean_num(total_ot_x):>12,.2f}")
    print(f"      Overtime Y:    EUR {clean_num(total_ot_y):>12,.2f}")
    print(f"    Service level:   {service_level*100:.1f}%  ({periods_no_bo}/{T})")
    print(f"    Fill rate:       {fill_rate*100:.2f}%")
    print(f"    Backorder units: {total_new_bo:,.0f}")
    print(f"    Safety stock:    {SAFETY_STOCK} units\n")

    # ══════════════════════════════════════════════════════════════════
    # Excel output
    # ══════════════════════════════════════════════════════════════════
    if os.path.exists(OUTPUT_FILE):
        wb = load_workbook(OUTPUT_FILE)
        if sheet in wb.sheetnames:
            del wb[sheet]
        ws = wb.create_sheet(sheet)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet

    NO_FILL    = PatternFill(fill_type=None)
    BLACK_FILL = PatternFill("solid", start_color="000000", end_color="000000")
    none_border = Border()
    bot_medium  = Border(bottom=Side(style="medium", color="000000"))
    bot_thin    = Border(bottom=Side(style="thin",   color="CCCCCC"))

    def plain(cell, val, bold=False, align="left", size=9, color="000000", fmt=None, border=None):
        cell.value = val
        cell.font  = Font(name="Calibri", bold=bold, size=size, color=color)
        cell.alignment = Alignment(horizontal=align, vertical="center")
        if border: cell.border = border
        if fmt:    cell.number_format = fmt

    def hdr(cell, val):
        cell.value = val
        cell.font  = Font(name="Calibri", size=9, color="FFFFFF")
        cell.fill  = BLACK_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def section_title(ws, r, last_col, text):
        plain(ws.cell(r, 2), text, size=8, color="888888", border=bot_medium)
        for col in range(3, last_col + 1):
            ws.cell(r, col).border = bot_medium

    ws.column_dimensions["A"].width = 1
    ws.column_dimensions["B"].width = 9
    for col in range(3, T + 4):
        ws.column_dimensions[get_column_letter(col)].width = 5.5
    last_col = T + 2

    # Title
    ws.merge_cells(f"B1:{get_column_letter(last_col)}1")
    plain(ws.cell(1, 2),
          f"Assignment 6a - Forecasted demand (safety stock, {label})",
          bold=True, size=13)

    # Cost summary
    summary_rows = [
        ("Setup cost",      setup_cost_plan),
        ("Holding cost",    total_holding),
        ("Backorder cost",  total_backorder),
        ("Investment X",    invest_X),
        ("Investment Y",    invest_Y),
        ("Overtime cost X", total_ot_x),
        ("Overtime cost Y", total_ot_y),
        ("Total cost",      total_cost),
    ]
    row = 3
    for lbl, val in summary_rows:
        plain(ws.cell(row, 2), lbl, size=9, color="555555")
        plain(ws.cell(row, 3), round(clean_num(val), 2),
              bold=(lbl == "Total cost"), fmt='"€"#,##0.00')
        ws.merge_cells(f"C{row}:{get_column_letter(last_col)}{row}")
        row += 1

    # Service metrics
    row += 1
    section_title(ws, row, last_col, f"Service metrics (end product, forecasted demand, {label})")
    row += 1
    plain(ws.cell(row, 2), "Service level", size=9, color="555555")
    plain(ws.cell(row, 3), f"{service_level*100:.1f}% ({periods_no_bo}/{T})")
    ws.merge_cells(f"C{row}:{get_column_letter(last_col)}{row}")
    row += 1
    plain(ws.cell(row, 2), "Fill rate", size=9, color="555555")
    plain(ws.cell(row, 3), f"{fill_rate*100:.2f}%")
    ws.merge_cells(f"C{row}:{get_column_letter(last_col)}{row}")
    row += 1
    plain(ws.cell(row, 2), "Total backorders", size=9, color="555555")
    plain(ws.cell(row, 3), f"{total_new_bo:,.0f} units",
          color="CC0000" if total_new_bo > 0 else "000000")
    ws.merge_cells(f"C{row}:{get_column_letter(last_col)}{row}")
    row += 1
    plain(ws.cell(row, 2), "Safety stock", size=9, color="555555")
    plain(ws.cell(row, 3), f"{SAFETY_STOCK} units (z={z_val}, {label})")
    ws.merge_cells(f"C{row}:{get_column_letter(last_col)}{row}")

    # Backorder schedule
    row += 2
    section_title(ws, row, last_col, f"Backorder schedule - {END_PRODUCT}")
    row += 1
    hdr(ws.cell(row, 2), "Part")
    for t in periods: hdr(ws.cell(row, t+2), str(t))
    row += 1
    plain(ws.cell(row, 2), END_PRODUCT, border=bot_thin)
    for t in periods:
        v = round(backorders[t])
        if v > 0:
            plain(ws.cell(row, t+2), v, bold=True, color="CC0000",
                  align="center", fmt='#,##0', border=bot_thin)
        else:
            plain(ws.cell(row, t+2), "", border=bot_thin)

    # Production schedule
    row += 2
    section_title(ws, row, last_col, "Production / order schedule (units)")
    row += 1
    hdr(ws.cell(row, 2), "Part")
    for t in periods: hdr(ws.cell(row, t+2), str(t))
    for i in PARTS:
        row += 1
        plain(ws.cell(row, 2), i, border=bot_thin)
        for t in periods:
            val = round(schedule[i][t]) if schedule[i][t] > 0.5 else ""
            plain(ws.cell(row, t+2), val, align="center", border=bot_thin)

    # Inventory levels
    row += 2
    section_title(ws, row, last_col, "Inventory levels (end of period)")
    row += 1
    hdr(ws.cell(row, 2), "Part")
    for t in periods: hdr(ws.cell(row, t+2), str(t))
    for i in PARTS:
        row += 1
        plain(ws.cell(row, 2), i, border=bot_thin)
        for t in periods:
            v = round(inventory[i][t])
            plain(ws.cell(row, t+2), v, align="center",
                  color="000000" if v > 0 else "BBBBBB",
                  fmt='#,##0', border=bot_thin)

    # Setup decisions
    row += 2
    section_title(ws, row, last_col, "Setup decisions")
    row += 1
    hdr(ws.cell(row, 2), "Part")
    for t in periods: hdr(ws.cell(row, t+2), str(t))
    for i in PARTS:
        row += 1
        plain(ws.cell(row, 2), i, border=bot_thin)
        for t in periods:
            plain(ws.cell(row, t+2),
                  "x" if schedule[i][t] > 0.5 else "",
                  align="center", border=bot_thin)

    wb.save(OUTPUT_FILE)
    print(f"  Written to {OUTPUT_FILE} -> sheet '{sheet}'")

# ══════════════════════════════════════════════════════════════════════════
# Comparison table: 95% vs 99%
# ══════════════════════════════════════════════════════════════════════════
if len(results) == 2:
    r95, r99 = results[0], results[1]
    print(f"\n{'='*74}")
    print(f"  COMPARISON (forecasted demand): {r95['label']} vs {r99['label']}")
    print(f"{'='*74}")
    print(f"{'Metric':<22}{'95% CSL':>14}{'99% CSL':>14}{'Diff':>14}{'Diff %':>10}")
    print(f"{'-'*74}")

    comp = [
        ("Safety stock (units)", r95["ss"],        r99["ss"]),
        ("Setup",                r95["setup"],     r99["setup"]),
        ("Holding",              r95["holding"],   r99["holding"]),
        ("Backorder",            r95["backorder"], r99["backorder"]),
        ("Investment X",         r95["invest_X"],  r99["invest_X"]),
        ("Investment Y",         r95["invest_Y"],  r99["invest_Y"]),
        ("Overtime X",           r95["ot_x"],      r99["ot_x"]),
        ("Overtime Y",           r95["ot_y"],      r99["ot_y"]),
        ("Total cost",           r95["total"],     r99["total"]),
    ]
    for lbl, v95, v99 in comp:
        diff = v99 - v95
        pct  = (diff / v95 * 100) if v95 != 0 else 0
        print(f"{lbl:<22}{v95:>14,.2f}{v99:>14,.2f}{diff:>+14,.2f}{pct:>+9.1f}%")

    print(f"")
    print(f"{'Service level (%)':<22}{r95['sl']*100:>14.1f}{r99['sl']*100:>14.1f}{(r99['sl']-r95['sl'])*100:>+14.1f}")
    print(f"{'Fill rate (%)':<22}{r95['fr']*100:>14.2f}{r99['fr']*100:>14.2f}{(r99['fr']-r95['fr'])*100:>+14.2f}")
    print(f"{'Backorder units':<22}{r95['bo_units']:>14,.0f}{r99['bo_units']:>14,.0f}{r99['bo_units']-r95['bo_units']:>+14,.0f}")
    print(f"{'='*74}\n")