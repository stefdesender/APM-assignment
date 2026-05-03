"""
Assignment 6b
Evaluate the 6a outsourcing plan under REALIZED demand.

"""

import gurobipy as gp
from gurobipy import GRB
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
from input_data import (
    PARTS, END_PRODUCT, T, BOM, LEAD_TIME, MIN_LOT,
    INIT_INV, SETUP_COST, HOLDING_COST,
    DEMAND_FORECAST, DEMAND_REALIZED, BACKORDER_COST
)


def get_parents(part):
    parents = {}
    for parent, children in BOM.items():
        if part in children:
            parents[parent] = children[part]
    return parents


def clean_num(val, tol=1e-6):
    v = float(val)
    return 0.0 if abs(v) < tol else round(v, 6)


#  Capacity / overtime / expansion (identical to 5a/5b) 
CAP_X_BASE     = 800
CAP_X_MAX_EXP  = 200
COST_EXP_X     = 10
CAP_X_OT_MAX   = 300
COST_OT_X      = 2

CAP_Y_BASE     = 60 * 24 * 7 - 80
CAP_Y_MAX_PCT  = 40
COST_EXP_Y_PCT = 1_500
CAP_Y_OT_MAX   = 38
COST_OT_Y      = 120

PROC_X = {END_PRODUCT: 1}
PROC_Y = {"B1401": 3, "B2302": 2}

#  Outsourcing parameters for B4702 
OUTSOURCED_PART      = "B4702"
DEFAULT_LEAD_TIME    = 1
OUTSOURCED_MIN_LOT   = 600
ORDER_COST_B4702     = 500   

periods = range(1, T + 1)


# Helper: solve 6a with given lead time, return the fixed plan
def solve_6a_plan(lead_time_b4702, order_cost=ORDER_COST_B4702, silent=True):
    """Solve the 6a model with a specific B4702 lead time and return the plan."""
    LEAD_TIME_LOC  = dict(LEAD_TIME);  LEAD_TIME_LOC[OUTSOURCED_PART]  = lead_time_b4702
    MIN_LOT_LOC    = dict(MIN_LOT);    MIN_LOT_LOC[OUTSOURCED_PART]    = OUTSOURCED_MIN_LOT
    SETUP_COST_LOC = dict(SETUP_COST); SETUP_COST_LOC[OUTSOURCED_PART] = 0

    m = gp.Model(f"APM_6a_lt{lead_time_b4702}")
    if silent:
        m.Params.OutputFlag = 0

    x      = m.addVars(PARTS, periods, name="x", vtype=GRB.INTEGER, lb=0)
    y      = m.addVars(PARTS, periods, name="y", vtype=GRB.BINARY)
    Iv     = m.addVars(PARTS, periods, name="I", vtype=GRB.INTEGER, lb=0)
    dx     = m.addVar(name="dx",     vtype=GRB.INTEGER, lb=0, ub=CAP_X_MAX_EXP)
    dy_pct = m.addVar(name="dy_pct", lb=0, ub=CAP_Y_MAX_PCT)
    ot_x   = m.addVars(periods, name="ot_x", vtype=GRB.INTEGER, lb=0, ub=CAP_X_OT_MAX)
    ot_y   = m.addVars(periods, name="ot_y", lb=0, ub=CAP_Y_OT_MAX)

    BIG_M = {i: sum(DEMAND_FORECAST) * 25 for i in PARTS}

    m.setObjective(
        gp.quicksum(SETUP_COST_LOC[i]*y[i,t] + HOLDING_COST[i]*Iv[i,t]
                    for i in PARTS for t in periods)
        + gp.quicksum(order_cost * y[OUTSOURCED_PART, t] for t in periods)
        + COST_EXP_X*dx + COST_EXP_Y_PCT*dy_pct
        + gp.quicksum(COST_OT_X*ot_x[t] + COST_OT_Y*ot_y[t] for t in periods),
        GRB.MINIMIZE)

    for i in PARTS:
        for t in periods:
            inv_prev = INIT_INV[i] if t == 1 else Iv[i, t-1]
            op = t - LEAD_TIME_LOC[i]
            rec = x[i, op] if op >= 1 else 0
            int_d = gp.quicksum(BOM[p][i]*x[p,t] for p in get_parents(i))
            ext_d = DEMAND_FORECAST[t-1] if i == END_PRODUCT else 0
            m.addConstr(Iv[i,t] == inv_prev + rec - int_d - ext_d)
            m.addConstr(x[i,t] >= MIN_LOT_LOC[i]*y[i,t])
            m.addConstr(x[i,t] <= BIG_M[i]*y[i,t])

    for t in periods:
        m.addConstr(gp.quicksum(PROC_X[i]*x[i,t] for i in PROC_X if i in PARTS)
                    <= CAP_X_BASE + dx + ot_x[t])
        m.addConstr(gp.quicksum(PROC_Y[i]*x[i,t] for i in PROC_Y if i in PARTS)
                    <= CAP_Y_BASE + (CAP_Y_BASE/100.0)*dy_pct + 60.0*ot_y[t])

    m.optimize()
    if m.status != GRB.OPTIMAL:
        return None

    return {
        "lead_time":   lead_time_b4702,
        "order_cost":  order_cost,
        "schedule":    {i: {t: x[i,t].X for t in periods} for i in PARTS},
        "y":           {i: {t: y[i,t].X for t in periods} for i in PARTS},
        "dx":          clean_num(dx.X),
        "dy_pct":      clean_num(dy_pct.X),
        "ot_x":        {t: clean_num(ot_x[t].X) for t in periods},
        "ot_y":        {t: clean_num(ot_y[t].X) for t in periods},
        "obj_6a":      m.ObjVal,
        "lead_time_table": LEAD_TIME_LOC,
        "setup_cost_table": SETUP_COST_LOC,
    }


# Helper: simulate 6a plan against realized demand 
def simulate_realized(plan):
    """Apply the frozen 6a plan to realized demand and compute all metrics."""
    schedule  = plan["schedule"]
    y_plan    = plan["y"]
    LT_LOC    = plan["lead_time_table"]
    SC_LOC    = plan["setup_cost_table"]

    inventory = {i: {} for i in PARTS}
    held      = {i: {} for i in PARTS}
    backorders = {t: 0.0 for t in periods}
    net_end    = {}

    
    for i in PARTS:
        if i == END_PRODUCT:
            continue
        inv_prev = float(INIT_INV[i])
        for t in periods:
            op = t - LT_LOC[i]
            rec = schedule[i].get(op, 0.0) if op >= 1 else 0.0
            int_d = sum(BOM[p][i]*schedule[p].get(t, 0.0) for p in get_parents(i))
            net = inv_prev + rec - int_d
            inventory[i][t] = clean_num(net)
            held[i][t]      = clean_num(max(0.0, net))
            inv_prev = net

   
    inv_prev = float(INIT_INV[END_PRODUCT])
    bo_prev  = 0.0
    for t in periods:
        op = t - LT_LOC[END_PRODUCT]
        rec = schedule[END_PRODUCT].get(op, 0.0) if op >= 1 else 0.0
        net = inv_prev - bo_prev + rec - DEMAND_REALIZED[t-1]
        net_end[t] = clean_num(net)
        if net >= 0:
            inventory[END_PRODUCT][t] = clean_num(net)
            held[END_PRODUCT][t]      = clean_num(net)
            backorders[t] = 0.0
        else:
            inventory[END_PRODUCT][t] = 0.0
            held[END_PRODUCT][t]      = 0.0
            backorders[t] = clean_num(-net)
        inv_prev = inventory[END_PRODUCT][t]
        bo_prev  = backorders[t]

    # Cost components
    setup_cost_fixed   = sum(SC_LOC[i] for i in PARTS for t in periods if schedule[i][t] > 0.5)
    n_orders_b4702     = sum(1 for t in periods if y_plan[OUTSOURCED_PART][t] > 0.5)
    order_cost_b4702   = plan["order_cost"] * n_orders_b4702
    total_holding      = sum(HOLDING_COST[i]*held[i][t] for i in PARTS for t in periods)
    total_backorder    = sum(BACKORDER_COST * backorders[t] for t in periods)
    invest_X           = COST_EXP_X     * plan["dx"]
    invest_Y           = COST_EXP_Y_PCT * plan["dy_pct"]
    total_ot_x         = sum(COST_OT_X * plan["ot_x"][t] for t in periods)
    total_ot_y         = sum(COST_OT_Y * plan["ot_y"][t] for t in periods)
    total_cost         = (setup_cost_fixed + order_cost_b4702 + total_holding
                          + total_backorder + invest_X + invest_Y
                          + total_ot_x + total_ot_y)

    # Service metrics 
    periods_no_bo = sum(1 for t in periods if backorders[t] == 0)
    service_level = periods_no_bo / T
    total_demand  = sum(DEMAND_REALIZED)
    total_new_bo  = 0.0
    prev_bo = 0.0
    for t in periods:
        new_bo = max(0.0, backorders[t] - prev_bo)
        total_new_bo += new_bo
        prev_bo = backorders[t]
    fill_rate     = 1.0 - (total_new_bo / total_demand) if total_demand > 0 else 1.0
    units_on_time = total_demand - total_new_bo

    return {
        "lead_time":         plan["lead_time"],
        "obj_6a":            plan["obj_6a"],
        "schedule":          schedule,
        "y_plan":            y_plan,
        "inventory":         inventory,
        "backorders":        backorders,
        "setup_cost":        setup_cost_fixed,
        "order_cost_b4702":  order_cost_b4702,
        "n_orders_b4702":    n_orders_b4702,
        "holding":           total_holding,
        "backorder_cost":    total_backorder,
        "invest_X":          invest_X,
        "invest_Y":          invest_Y,
        "ot_x":              total_ot_x,
        "ot_y":              total_ot_y,
        "total_cost":        total_cost,
        "service_level":     service_level,
        "periods_no_bo":     periods_no_bo,
        "fill_rate":         fill_rate,
        "units_on_time":     units_on_time,
        "total_new_bo":      total_new_bo,
        "total_demand":      total_demand,
        "ot_x_vals":         plan["ot_x"],
        "ot_y_vals":         plan["ot_y"],
        "dx":                plan["dx"],
        "dy_pct":            plan["dy_pct"],
    }


# Main: run the base 6b case (lead time 1) + lead-time sensitivity
print("="*60)
print("Assignment 6b - base case (lead time 1)")
print("="*60)
plan_main = solve_6a_plan(DEFAULT_LEAD_TIME, silent=False)
res_main  = simulate_realized(plan_main)

# Lead-time sensitivity
LEAD_TIME_SWEEP = [1, 2, 3, 4]
print("\n" + "="*60)
print("Assignment 6b - lead-time sensitivity (realized demand)")
print("="*60)
sens_results = []
for lt in LEAD_TIME_SWEEP:
    print(f"  Solving + simulating with lead time {lt} week(s) ...")
    plan_lt = solve_6a_plan(lt, silent=True)
    if plan_lt:
        res_lt = simulate_realized(plan_lt)
        sens_results.append(res_lt)
        print(f"    -> Total: €{res_lt['total_cost']:>10,.2f}  | "
              f"SL: {res_lt['service_level']*100:>5.1f}%  | "
              f"FR: {res_lt['fill_rate']*100:>6.2f}%  | "
              f"BO units: {res_lt['total_new_bo']:>5.0f}")


# Excel output
OUTPUT_FILE = "OUTPUT.xlsx"

if os.path.exists(OUTPUT_FILE):
    wb = load_workbook(OUTPUT_FILE)
else:
    wb = Workbook()
    wb.remove(wb.active)

#  Styles 
NO_FILL    = PatternFill(fill_type=None)
BLACK_FILL = PatternFill("solid", start_color="000000", end_color="000000")
none_border = Border()
bot_medium  = Border(bottom=Side(style="medium", color="000000"))
bot_thin    = Border(bottom=Side(style="thin",   color="CCCCCC"))


def plain(cell, val, bold=False, align="left", size=9, color="000000", fmt=None, border=None):
    cell.value = val
    cell.font  = Font(name="Calibri", bold=bold, size=size, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.fill  = NO_FILL
    if border: cell.border = border
    if fmt:    cell.number_format = fmt


def hdr(cell, val):
    cell.value = val
    cell.font  = Font(name="Calibri", bold=False, size=9, color="FFFFFF")
    cell.fill  = BLACK_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = none_border


def section_title(ws, r, last_col, text):
    ws.row_dimensions[r].height = 14
    plain(ws.cell(r, 2), text, size=8, color="888888", border=bot_medium)
    for col in range(3, last_col + 1):
        ws.cell(r, col).border = bot_medium


# Sheet: Output_6b 
SHEET_MAIN = "Output_6b"
if SHEET_MAIN in wb.sheetnames:
    del wb[SHEET_MAIN]
ws = wb.create_sheet(SHEET_MAIN)

ws.column_dimensions["A"].width = 1
ws.column_dimensions["B"].width = 9
for col in range(3, T + 4):
    ws.column_dimensions[get_column_letter(col)].width = 5.5
last_col = T + 2

# Title
ws.row_dimensions[1].height = 26
ws.merge_cells(f"B1:{get_column_letter(last_col)}1")
plain(ws.cell(1, 2),
      f"Assignment 6b - Realized demand evaluation (outsourcing B4702, order cost €{ORDER_COST_B4702}, lead time {DEFAULT_LEAD_TIME})",
      bold=True, size=13)

# Cost summary
ws.row_dimensions[2].height = 4
cost_rows = [
    ("Setup cost",          res_main["setup_cost"],         False),
    ("Order cost B4702",    res_main["order_cost_b4702"],   False),
    ("Holding cost",        res_main["holding"],            False),
    ("Backorder cost",      res_main["backorder_cost"],     False),
    ("Investment cost X",   res_main["invest_X"],           False),
    ("Investment cost Y",   res_main["invest_Y"],           False),
    ("Overtime cost X",     res_main["ot_x"],               False),
    ("Overtime cost Y",     res_main["ot_y"],               False),
    ("Total cost (6b)",     res_main["total_cost"],         True),
    ("Total cost (6a)",     res_main["obj_6a"],             False),
    ("Difference",          res_main["total_cost"] - res_main["obj_6a"], False),
]
for r, (label, val, bold) in enumerate(cost_rows, start=3):
    ws.row_dimensions[r].height = 16
    plain(ws.cell(r, 2), label, size=9, color="555555")
    fmt = '€ #,##0.00;-€ #,##0.00' if label == "Difference" else '€ #,##0.00'
    plain(ws.cell(r, 3), round(clean_num(val), 2), bold=bold, size=9, fmt=fmt)
    ws.merge_cells(f"C{r}:{get_column_letter(last_col)}{r}")

# Service metrics
ws.row_dimensions[14].height = 6
ws.row_dimensions[15].height = 14
section_title(ws, 15, last_col, "Service metrics (end product)")
for r, (label, text, red) in enumerate([
    ("Service level",   f"{res_main['service_level']*100:.1f}%  ({res_main['periods_no_bo']}/{T} periods without backorder)", False),
    ("Fill rate",       f"{res_main['fill_rate']*100:.2f}%  ({res_main['units_on_time']:,.0f} / {res_main['total_demand']:,.0f} units on time)", False),
    ("Total backorder", f"{res_main['total_new_bo']:,.0f} units", res_main['total_new_bo'] > 0),
    ("# B4702 orders",  f"{res_main['n_orders_b4702']} orders", False),
], start=16):
    ws.row_dimensions[r].height = 16
    plain(ws.cell(r, 2), label, size=9, color="555555")
    plain(ws.cell(r, 3), text, size=9, color="CC0000" if red else "000000")
    ws.merge_cells(f"C{r}:{get_column_letter(last_col)}{r}")

# Backorder schedule
r = 21
section_title(ws, r, last_col, f"Backorder schedule - {END_PRODUCT} (end product only)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "Part")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))
r += 1
ws.row_dimensions[r].height = 15
plain(ws.cell(r, 2), END_PRODUCT, size=9, border=bot_thin)
for t in periods:
    v = round(res_main["backorders"][t])
    if v > 0:
        plain(ws.cell(r, t+2), v, bold=True, size=9, align="center",
              color="CC0000", fmt='#,##0', border=bot_thin)
    else:
        plain(ws.cell(r, t+2), "", size=9, align="center", border=bot_thin)

# Demand vs delivered
r += 2
section_title(ws, r, last_col, "Demand vs delivered (end product)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))
r += 1
ws.row_dimensions[r].height = 15
plain(ws.cell(r, 2), "Demand", size=9, border=bot_thin)
for t in periods:
    plain(ws.cell(r, t+2), DEMAND_REALIZED[t-1], size=9, align="center",
          fmt='#,##0', border=bot_thin)
r += 1
ws.row_dimensions[r].height = 15
plain(ws.cell(r, 2), "Delivered", size=9, border=bot_thin)
prev_bo = 0.0
for t in periods:
    new_bo_t = max(0.0, res_main["backorders"][t] - prev_bo)
    delivered = DEMAND_REALIZED[t-1] - new_bo_t
    shortage  = delivered < DEMAND_REALIZED[t-1]
    plain(ws.cell(r, t+2), round(delivered), size=9, align="center",
          color="CC0000" if shortage else "000000",
          fmt='#,##0', border=bot_thin)
    prev_bo = res_main["backorders"][t]

# Production / order schedule
r += 2
section_title(ws, r, last_col, "Production / order schedule (units, from 6a plan)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "Part")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))
for i in PARTS:
    r += 1
    ws.row_dimensions[r].height = 15
    label_i = i + " *" if i == OUTSOURCED_PART else i
    plain(ws.cell(r, 2), label_i, size=9, border=bot_thin,
          bold=(i == OUTSOURCED_PART))
    for t in periods:
        val = round(res_main["schedule"][i][t]) if res_main["schedule"][i][t] > 0.5 else ""
        plain(ws.cell(r, t+2), val, bold=bool(val), size=9, align="center",
              fmt='#,##0' if val != "" else None, border=bot_thin)

# Inventory levels (realized)
r += 2
section_title(ws, r, last_col, "Inventory levels (end of period, realized)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "Part")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))
for i in PARTS:
    r += 1
    ws.row_dimensions[r].height = 15
    plain(ws.cell(r, 2), i, size=9, border=bot_thin)
    for t in periods:
        v = round(res_main["inventory"][i][t])
        plain(ws.cell(r, t+2), v, size=9, align="center",
              color="BBBBBB" if v == 0 else "000000",
              fmt='#,##0', border=bot_thin)

# Overtime usage
r += 2
section_title(ws, r, last_col, "Overtime usage (from 6a plan, unchanged)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))
for label, vals, fmt in [
    ("X  (units OT)", res_main["ot_x_vals"], '#,##0'),
    ("Y  (hours OT)", res_main["ot_y_vals"], '0.00'),
]:
    r += 1
    ws.row_dimensions[r].height = 15
    plain(ws.cell(r, 2), label, size=9, border=bot_thin)
    for t in periods:
        v = round(vals[t], 2)
        col = "CC0000" if v else "BBBBBB"
        plain(ws.cell(r, t+2), v if v else "", bold=bool(v), size=9,
              align="center", color=col,
              fmt=fmt if v else None, border=bot_thin)

# Setup / order decisions
r += 2
section_title(ws, r, last_col, "Setup / order decisions  (* = outsourced order)")
r += 1
ws.row_dimensions[r].height = 16
hdr(ws.cell(r, 2), "Part")
for t in periods:
    hdr(ws.cell(r, t+2), str(t))
for i in PARTS:
    r += 1
    ws.row_dimensions[r].height = 15
    label_i = i + " *" if i == OUTSOURCED_PART else i
    plain(ws.cell(r, 2), label_i, size=9, border=bot_thin,
          bold=(i == OUTSOURCED_PART))
    for t in periods:
        plain(ws.cell(r, t+2), "x" if res_main["schedule"][i][t] > 0.5 else "",
              bold=True, size=9, align="center", border=bot_thin)


# Sheet: Output_6b_sens_LT (lead-time sensitivity under realized demand)
SHEET_LT = "Output_6b_sens_LT"
if SHEET_LT in wb.sheetnames:
    del wb[SHEET_LT]
wsl = wb.create_sheet(SHEET_LT)

wsl.column_dimensions["A"].width = 1
wsl.column_dimensions["B"].width = 24
for col in range(3, 3 + len(sens_results)):
    wsl.column_dimensions[get_column_letter(col)].width = 14

# Title
wsl.row_dimensions[1].height = 26
last_col_l = 2 + len(sens_results)
wsl.merge_cells(f"B1:{get_column_letter(last_col_l)}1")
plain(wsl.cell(1, 2),
      f"Assignment 6b - Lead-time sensitivity B4702 (order cost €{ORDER_COST_B4702}, realized demand)",
      bold=True, size=13)

wsl.row_dimensions[2].height = 4

hdr(wsl.cell(3, 2), "Lead time (weeks)")
for j, res in enumerate(sens_results):
    hdr(wsl.cell(3, 3 + j), f"{res['lead_time']} wk")

lt_base = sens_results[0]
metrics_lt = [
    ("Total cost 6b (EUR)",      "total_cost",      '"€"#,##0.00',  True),
    ("  vs LT=1 (best case)",    "_diff_base",      '"€"#,##0.00;[Red]-"€"#,##0.00', False),
    ("Total cost 6a (EUR)",      "obj_6a",          '"€"#,##0.00',  False),
    ("Setup cost (EUR)",         "setup_cost",      '"€"#,##0.00',  False),
    ("Order cost B4702 (EUR)",   "order_cost_b4702",'"€"#,##0.00',  False),
    ("Holding cost (EUR)",       "holding",         '"€"#,##0.00',  False),
    ("Backorder cost (EUR)",     "backorder_cost",  '"€"#,##0.00',  False),
    ("Investment X (EUR)",       "invest_X",        '"€"#,##0.00',  False),
    ("Investment Y (EUR)",       "invest_Y",        '"€"#,##0.00',  False),
    ("Overtime X (EUR)",         "ot_x",            '"€"#,##0.00',  False),
    ("Overtime Y (EUR)",         "ot_y",            '"€"#,##0.00',  False),
    ("Service level (%)",        "_sl_pct",         '0.0',          True),
    ("Fill rate (%)",            "_fr_pct",         '0.00',         True),
    ("Backorder units",          "total_new_bo",    '#,##0',        False),
    ("Periods w/o backorder",    "periods_no_bo",   '0',            False),
    ("# B4702 orders",           "n_orders_b4702",  '0',            False),
    ("Permanent exp. X",         "dx",              '#,##0',        False),
    ("Permanent exp. Y (%)",     "dy_pct",          '0.0',          False),
]
for i, (label, key, fmt, bold) in enumerate(metrics_lt):
    r = 4 + i
    wsl.row_dimensions[r].height = 16
    plain(wsl.cell(r, 2), label, size=9, bold=bold,
          color="000000" if bold or not label.startswith("  ") else "777777")
    for j, res in enumerate(sens_results):
        if key == "_diff_base":
            v = res["total_cost"] - lt_base["total_cost"]
            color = "009900" if v < 0 else ("CC0000" if v > 0 else "777777")
            plain(wsl.cell(r, 3 + j), round(v, 2), size=9, align="right",
                  fmt=fmt, color=color, bold=bold)
        elif key == "_sl_pct":
            plain(wsl.cell(r, 3 + j), round(res["service_level"]*100, 2),
                  size=9, align="right", fmt=fmt, bold=bold)
        elif key == "_fr_pct":
            plain(wsl.cell(r, 3 + j), round(res["fill_rate"]*100, 2),
                  size=9, align="right", fmt=fmt, bold=bold)
        else:
            plain(wsl.cell(r, 3 + j), round(res[key], 2), size=9,
                  align="right", fmt=fmt, bold=bold)


wb.save(OUTPUT_FILE)
print(f"\n{'='*60}")
print(f"Results written to {OUTPUT_FILE}")
print(f"  -> sheet '{SHEET_MAIN}'  (main 6b case)")
print(f"  -> sheet '{SHEET_LT}'  (lead-time sensitivity)")
print(f"{'='*60}")

print(f"\nMain 6b case (lead time {DEFAULT_LEAD_TIME}, order cost €{ORDER_COST_B4702}):")
print(f"  Total cost (6b): EUR {res_main['total_cost']:,.2f}")
print(f"  Total cost (6a): EUR {res_main['obj_6a']:,.2f}")
print(f"  Difference:      EUR {res_main['total_cost'] - res_main['obj_6a']:+,.2f}")
print(f"  Service level:   {res_main['service_level']*100:.1f}%")
print(f"  Fill rate:       {res_main['fill_rate']*100:.2f}%")
print(f"  Backorder units: {res_main['total_new_bo']:.0f}")

print(f"\nLead-time sensitivity (realized demand):")
print(f"{'LT':>4} | {'Total 6b':>13} | {'vs LT=1':>12} | {'SL':>6} | {'FR':>7} | {'BO units':>9}")
print("-" * 75)
for res in sens_results:
    diff = res["total_cost"] - lt_base["total_cost"]
    print(f"{res['lead_time']:>3}wk | €{res['total_cost']:>12,.2f} | "
          f"€{diff:>+11,.2f} | "
          f"{res['service_level']*100:>5.1f}% | "
          f"{res['fill_rate']*100:>6.2f}% | "
          f"{res['total_new_bo']:>9.0f}")