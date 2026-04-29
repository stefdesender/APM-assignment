"""
APM Project 2026 - Assignment 6a (forecast with outsourcing)
"""

import gurobipy as gp
from gurobipy import GRB
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
import os
from input_data import (
    PARTS, END_PRODUCT, T, BOM, LEAD_TIME, MIN_LOT,
    INIT_INV, SETUP_COST, HOLDING_COST, DEMAND_FORECAST
)

# ── OUTSOURCING SETTINGS ───────────────────────────────
OUTSOURCED_PART = "B4702"
ORDER_COST = 332105.08

def get_parents(part):
    parents = {}
    for parent, children in BOM.items():
        if part in children:
            parents[parent] = children[part]
    return parents

def clean_num(val, tol=1e-6):
    v = float(val)
    return 0.0 if abs(v) < tol else round(v, 6)

# ── COPY & MODIFY PARAMETERS ───────────────────────────
LEAD_TIME_6 = LEAD_TIME.copy()
SETUP_COST_6 = SETUP_COST.copy()

LEAD_TIME_6[OUTSOURCED_PART] = 1
SETUP_COST_6[OUTSOURCED_PART] = ORDER_COST

# ── CAPACITY PARAMETERS (zelfde als 5a) ─────────────────
CAP_X_BASE     = 800
CAP_X_MAX_EXP  = 200
COST_EXP_X     = 10
CAP_X_OT_MAX   = 300
COST_OT_X      = 2

CAP_Y_BASE     = 60 * 24 * 7 - 80
CAP_Y_MAX_PCT  = 40
COST_EXP_Y_PCT = 1500
CAP_Y_OT_MAX   = 38
COST_OT_Y      = 120

PROC_X = {END_PRODUCT: 1}
PROC_Y = {"B1401": 3, "B2302": 2}

# ── MODEL ──────────────────────────────────────────────
model = gp.Model("APM_Assignment6a")

periods = range(1, T + 1)
parts   = PARTS

x      = model.addVars(parts, periods, name="x", vtype=GRB.INTEGER, lb=0)
y      = model.addVars(parts, periods, name="y", vtype=GRB.BINARY)
I      = model.addVars(parts, periods, name="I", vtype=GRB.INTEGER, lb=0)
dx     = model.addVar(name="dx", vtype=GRB.INTEGER, lb=0, ub=CAP_X_MAX_EXP)
dy_pct = model.addVar(name="dy_pct", lb=0, ub=CAP_Y_MAX_PCT)
ot_x   = model.addVars(periods, name="ot_x", vtype=GRB.INTEGER, lb=0, ub=CAP_X_OT_MAX)
ot_y   = model.addVars(periods, name="ot_y", lb=0, ub=CAP_Y_OT_MAX)

BIG_M = {i: sum(DEMAND_FORECAST) * 25 for i in parts}

# ── OBJECTIVE ──────────────────────────────────────────
model.setObjective(
    gp.quicksum(SETUP_COST_6[i]*y[i,t] + HOLDING_COST[i]*I[i,t]
                for i in parts for t in periods)
    + COST_EXP_X*dx + COST_EXP_Y_PCT*dy_pct
    + gp.quicksum(COST_OT_X*ot_x[t] + COST_OT_Y*ot_y[t] for t in periods),
    GRB.MINIMIZE)

# ── CONSTRAINTS ────────────────────────────────────────
for i in parts:
    for t in periods:
        inv_prev = INIT_INV[i] if t == 1 else I[i, t-1]
        op = t - LEAD_TIME_6[i]
        rec = x[i, op] if op >= 1 else 0
        int_d = gp.quicksum(BOM[p][i]*x[p,t] for p in get_parents(i))
        ext_d = DEMAND_FORECAST[t-1] if i == END_PRODUCT else 0

        model.addConstr(I[i,t] == inv_prev + rec - int_d - ext_d)
        model.addConstr(x[i,t] >= MIN_LOT[i]*y[i,t])
        model.addConstr(x[i,t] <= BIG_M[i]*y[i,t])

for t in periods:
    model.addConstr(
        gp.quicksum(PROC_X[i]*x[i,t] for i in PROC_X if i in parts)
        <= CAP_X_BASE + dx + ot_x[t])
    model.addConstr(
        gp.quicksum(PROC_Y[i]*x[i,t] for i in PROC_Y if i in parts)
        <= CAP_Y_BASE + (CAP_Y_BASE/100.0)*dy_pct + 60.0*ot_y[t])

model.optimize()

# ── RESULTS ────────────────────────────────────────────
if model.status == GRB.OPTIMAL:

    total_setup   = sum(SETUP_COST_6[i]*y[i,t].X for i in parts for t in periods)
    total_holding = sum(HOLDING_COST[i]*I[i,t].X for i in parts for t in periods)

    # ── JSON OUTPUT ─────────────────────────────────────
    output = {
        "status": "OPTIMAL",
        "total_cost": model.ObjVal,
        "outsourcing": {
            "part": OUTSOURCED_PART,
            "order_cost": ORDER_COST,
            "lead_time_new": LEAD_TIME_6[OUTSOURCED_PART]
        },
        "cost_breakdown": {
            "setup_cost": total_setup,
            "holding_cost": total_holding
        },
        "production_schedule": {
            i: {t: x[i,t].X for t in periods if x[i,t].X > 0.5}
            for i in parts
        }
    }

    with open("output_6_1.json", "w") as f:
        json.dump(output, f, indent=2)