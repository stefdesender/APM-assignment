"""
Assignment 6a
Idea: Sensitivity analysis on the permanent expansion of Workstation X.

"""

import gurobipy as gp
from gurobipy import GRB
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference, Series
import os

from input_data import (
    PARTS, END_PRODUCT, T, BOM, LEAD_TIME, MIN_LOT,
    INIT_INV, SETUP_COST, HOLDING_COST,
    DEMAND_FORECAST, DEMAND_REALIZED, BACKORDER_COST
)

#  helpers 
def get_parents(part):
    parents = {}
    for parent, children in BOM.items():
        if part in children:
            parents[parent] = children[part]
    return parents

def clean_num(val, tol=1e-6):
    v = float(val)
    return 0.0 if abs(v) < tol else round(v, 6)

#  constants (identical to 5a/5b) 
CAP_X_BASE     = 800
COST_EXP_X     = 10
CAP_X_OT_MAX   = 300
COST_OT_X      = 2

CAP_Y_BASE     = 60 * 24 * 7 - 80          
COST_EXP_Y_PCT = 1_500
CAP_Y_OT_MAX   = 38
COST_OT_Y      = 120

# dy_pct fixed at 5a optimal throughout
DY_FIXED = 14.1

PROC_X = {END_PRODUCT: 1}
PROC_Y = {"B1401": 3, "B2302": 2}

periods = range(1, T + 1)
BIG_M   = {i: sum(DEMAND_FORECAST) * 25 for i in PARTS}

#  sensitivity range 

REF_DX     = 200
STEP       = 50
CAP_VALUES = (
    [REF_DX - 2*STEP, REF_DX - 1*STEP]
    + [REF_DX + i*STEP for i in range(0, 9)]
)


#  solve one scenario 
def solve_scenario(dx_fixed: int) -> dict | None:
    m = gp.Model(f"APM_6a_dx{dx_fixed}")
    m.setParam("OutputFlag", 0)

    x      = m.addVars(PARTS, periods, name="x",  vtype=GRB.INTEGER, lb=0)
    y      = m.addVars(PARTS, periods, name="y",  vtype=GRB.BINARY)
    I      = m.addVars(PARTS, periods, name="I",  vtype=GRB.INTEGER, lb=0)
    dx     = m.addVar(name="dx",     vtype=GRB.INTEGER, lb=dx_fixed,  ub=dx_fixed)
    dy_pct = m.addVar(name="dy_pct", lb=DY_FIXED, ub=DY_FIXED)
    ot_x   = m.addVars(periods, name="ot_x", vtype=GRB.INTEGER, lb=0, ub=CAP_X_OT_MAX)
    ot_y   = m.addVars(periods, name="ot_y", lb=0, ub=CAP_Y_OT_MAX)

    m.setObjective(
        gp.quicksum(SETUP_COST[i]*y[i,t] + HOLDING_COST[i]*I[i,t]
                    for i in PARTS for t in periods)
        + COST_EXP_X * dx + COST_EXP_Y_PCT * dy_pct
        + gp.quicksum(COST_OT_X*ot_x[t] + COST_OT_Y*ot_y[t] for t in periods),
        GRB.MINIMIZE
    )

    for i in PARTS:
        for t in periods:
            inv_prev = INIT_INV[i] if t == 1 else I[i, t-1]
            op  = t - LEAD_TIME[i]
            rec = x[i, op] if op >= 1 else 0
            int_d = gp.quicksum(BOM[p][i]*x[p,t] for p in get_parents(i))
            ext_d = DEMAND_FORECAST[t-1] if i == END_PRODUCT else 0
            m.addConstr(I[i,t] == inv_prev + rec - int_d - ext_d)
            m.addConstr(x[i,t] >= MIN_LOT[i] * y[i,t])
            m.addConstr(x[i,t] <= BIG_M[i]   * y[i,t])

    for t in periods:
        m.addConstr(
            gp.quicksum(PROC_X.get(i,0)*x[i,t] for i in PARTS)
            <= CAP_X_BASE + dx + ot_x[t]
        )
        m.addConstr(
            gp.quicksum(PROC_Y.get(i,0)*x[i,t] for i in PARTS)
            <= CAP_Y_BASE + (CAP_Y_BASE / 100.0) * dy_pct + 60.0 * ot_y[t]
        )

    m.optimize()
    if m.status != GRB.OPTIMAL:
        return None

    return dict(
        dx_fixed         = dx_fixed,
        invest_X         = COST_EXP_X     * dx_fixed,
        invest_Y         = COST_EXP_Y_PCT * DY_FIXED,
        forecast_setup   = sum(SETUP_COST[i]   * y[i,t].X for i in PARTS for t in periods),
        forecast_holding = sum(HOLDING_COST[i] * I[i,t].X for i in PARTS for t in periods),
        forecast_ot_x    = sum(COST_OT_X * ot_x[t].X for t in periods),
        forecast_ot_y    = sum(COST_OT_Y * ot_y[t].X for t in periods),
        forecast_total   = m.ObjVal,
        schedule         = {i: {t: x[i,t].X   for t in periods} for i in PARTS},
        ot_x_vals        = {t: clean_num(ot_x[t].X) for t in periods},
        ot_y_vals        = {t: clean_num(ot_y[t].X) for t in periods},
    )


#  simulate realized demand 
def simulate_realized(res: dict) -> dict:
    schedule = res["schedule"]

    setup_cost_fixed = sum(
        SETUP_COST[i]
        for i in PARTS for t in periods
        if schedule[i].get(t, 0.0) > 0.5
    )

    inventory  = {i: {} for i in PARTS}
    held       = {i: {} for i in PARTS}
    backorders = {t: 0.0 for t in periods}

    for i in PARTS:
        if i == END_PRODUCT:
            continue
        inv_prev = float(INIT_INV[i])
        for t in periods:
            op    = t - LEAD_TIME[i]
            rec   = schedule[i].get(op, 0.0) if op >= 1 else 0.0
            int_d = sum(BOM[p][i] * schedule[p].get(t, 0.0) for p in get_parents(i))
            net   = inv_prev + rec - int_d
            inventory[i][t] = clean_num(net)
            held[i][t]      = clean_num(max(0.0, net))
            inv_prev = net

    inv_prev = float(INIT_INV[END_PRODUCT])
    bo_prev  = 0.0
    for t in periods:
        op  = t - LEAD_TIME[END_PRODUCT]
        rec = schedule[END_PRODUCT].get(op, 0.0) if op >= 1 else 0.0
        net = inv_prev - bo_prev + rec - DEMAND_REALIZED[t-1]
        if net >= 0:
            inventory[END_PRODUCT][t] = clean_num(net)
            held[END_PRODUCT][t]      = clean_num(net)
            backorders[t] = 0.0
        else:
            inventory[END_PRODUCT][t] = 0.0
            held[END_PRODUCT][t]      = 0.0
            backorders[t] = clean_num(-net)
        inv_prev = inventory[END_PRODUCT][t]
        bo_prev  = backorders[t]

    total_holding   = sum(HOLDING_COST[i] * held[i][t] for i in PARTS for t in periods)
    total_backorder = sum(BACKORDER_COST  * backorders[t] for t in periods)
    total_cost      = (
        setup_cost_fixed + total_holding + total_backorder
        + res["invest_X"] + res["invest_Y"]
        + sum(COST_OT_X * res["ot_x_vals"][t] for t in periods)
        + sum(COST_OT_Y * res["ot_y_vals"][t] for t in periods)
    )

    periods_no_bo = sum(1 for t in periods if backorders[t] == 0)
    total_demand  = sum(DEMAND_REALIZED)
    total_new_bo  = 0.0
    prev_bo = 0.0
    for t in periods:
        new_bo = max(0.0, backorders[t] - prev_bo)
        total_new_bo += new_bo
        prev_bo = backorders[t]

    return dict(
        setup_cost     = setup_cost_fixed,
        holding_cost   = total_holding,
        backorder_cost = total_backorder,
        total_cost     = total_cost,
        service_level  = periods_no_bo / T,
        fill_rate      = 1.0 - (total_new_bo / total_demand) if total_demand > 0 else 1.0,
        total_new_bo   = total_new_bo,
        backorders     = backorders,
    )


#  run all scenarios 
print("Running sensitivity analysis on permanent expansion of Workstation X ...")
print(f"  dy_pct fixed at {DY_FIXED}% (5a optimal) for all scenarios")
print(f"{'dx (units)':>12}  {'Cost forecast':>14}  {'Cost realized':>14}  "
      f"{'SL':>6}  {'FR':>7}  {'Backorder':>12}")
print("-" * 76)

results = []
for dx_val in CAP_VALUES:
    res = solve_scenario(dx_val)
    if res is None:
        print(f"{dx_val:>11}u  INFEASIBLE")
        continue
    sim = simulate_realized(res)
    res["sim"] = sim
    results.append(res)
    print(f"{dx_val:>11}u  "
          f"€ {res['forecast_total']:>12,.2f}  "
          f"€ {sim['total_cost']:>12,.2f}  "
          f"{sim['service_level']*100:>5.1f}%  "
          f"{sim['fill_rate']*100:>6.2f}%  "
          f"€ {sim['backorder_cost']:>10,.2f}")


#  Excel output 
OUTPUT_FILE = "OUTPUT.xlsx"
SHEET_NAME  = "Output_6a_maxX"

if os.path.exists(OUTPUT_FILE):
    wb = load_workbook(OUTPUT_FILE)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)
else:
    wb = Workbook()
    ws  = wb.active
    ws.title = SHEET_NAME

NO_FILL    = PatternFill(fill_type=None)
BLACK_FILL = PatternFill("solid", start_color="000000", end_color="000000")
LGREY_FILL = PatternFill("solid", start_color="D9D9D9", end_color="D9D9D9")
GREEN_FILL = PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA")
none_border = Border()
bot_medium  = Border(bottom=Side(style="medium", color="000000"))
full_thin   = Border(
    left  =Side(style="thin", color="CCCCCC"),
    right =Side(style="thin", color="CCCCCC"),
    top   =Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

def plain(cell, val, bold=False, align="left", size=9,
          color="000000", fmt=None, border=None, fill=None):
    cell.value = val
    cell.font  = Font(name="Calibri", bold=bold, size=size, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.fill  = fill if fill else NO_FILL
    if border: cell.border = border
    if fmt:    cell.number_format = fmt

def hdr(cell, val):
    cell.value = val
    cell.font  = Font(name="Calibri", bold=True, size=9, color="FFFFFF")
    cell.fill  = BLACK_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = none_border

def section_title(ws, r, last_col, text):
    ws.row_dimensions[r].height = 14
    plain(ws.cell(r, 2), text, size=8, color="888888", border=bot_medium)
    for col in range(3, last_col + 1):
        ws.cell(r, col).border = bot_medium

ws.column_dimensions["A"].width = 1
for col_letter, w in [("B",26),("C",12),("D",12),("E",15),("F",15),
                       ("G",15),("H",15),("I",11),("J",11)]:
    ws.column_dimensions[col_letter].width = w
LAST_COL = 10

# Title
ws.row_dimensions[1].height = 28
ws.merge_cells("B1:J1")
plain(ws.cell(1,2),
      "Assignment 6a - Sensitivity analysis: permanent expansion Workstation X",
      bold=True, size=13)
ws.row_dimensions[2].height = 14
ws.merge_cells("B2:J2")
plain(ws.cell(2,2),
      f"dx swept around 5a optimal (200 u) in steps of {STEP} u  |  "
      f"dy_pct fixed at {DY_FIXED}% (5a optimal) throughout.",
      size=8, color="555555")

best_realized = min(res["sim"]["total_cost"] for res in results)

#  Summary table 
r = 4
section_title(ws, r, LAST_COL,
              "Sensitivity table: forecast cost vs realized cost per dx scenario")
r += 1
ws.row_dimensions[r].height = 18
for c, h in enumerate([
    "dx (units)", "Invest X (€)",
    "Forecast total (€)", "Realized total (€)",
    "Holding real. (€)", "Backorder real. (€)",
    "Service level", "Fill rate"
], start=2):
    hdr(ws.cell(r, c), h)

for res in results:
    r += 1
    ws.row_dimensions[r].height = 16
    sim     = res["sim"]
    is_ref  = (res["dx_fixed"] == REF_DX)
    is_best = abs(sim["total_cost"] - best_realized) < 0.01
    row_fill = LGREY_FILL if is_ref else (GREEN_FILL if is_best else NO_FILL)
    suffix   = " <- ref (5a)" if is_ref else (" <- best" if is_best else "")

    for c, (val, bold, align, fmt, color) in enumerate([
        (f"{res['dx_fixed']} u{suffix}",        False, "left",   None,           "000000"),
        (res["invest_X"],                        False, "right",  '"€"#,##0.00',"000000"),
        (res["forecast_total"],                  False, "right",  '"€"#,##0.00',"000000"),
        (sim["total_cost"],                      is_best,"right", '"€"#,##0.00',"000000"),
        (sim["holding_cost"],                    False, "right",  '"€"#,##0.00',"000000"),
        (sim["backorder_cost"],                  False, "right",  '"€"#,##0.00',
         "CC0000" if sim["backorder_cost"] > 0 else "000000"),
        (f"{sim['service_level']*100:.1f}%",     False, "center", None,           "000000"),
        (f"{sim['fill_rate']*100:.2f}%",         False, "center", None,           "000000"),
    ], start=2):
        plain(ws.cell(r, c), val, bold=bold, align=align,
              fmt=fmt, color=color, border=full_thin, fill=row_fill)

#  Full cost breakdown 
r += 2
section_title(ws, r, LAST_COL, "Full cost breakdown (realized demand)")
r += 1
ws.row_dimensions[r].height = 18
for c, h in enumerate([
    "dx (units)", "Setup", "Holding", "Backorder",
    "OT cost X", "OT cost Y", "Invest X", "Invest Y", "Total"
], start=2):
    hdr(ws.cell(r, c), h)

for res in results:
    r += 1
    ws.row_dimensions[r].height = 16
    sim        = res["sim"]
    ot_x_total = sum(COST_OT_X * res["ot_x_vals"][t] for t in periods)
    ot_y_total = sum(COST_OT_Y * res["ot_y_vals"][t] for t in periods)
    is_ref  = (res["dx_fixed"] == REF_DX)
    is_best = abs(sim["total_cost"] - best_realized) < 0.01
    row_fill = LGREY_FILL if is_ref else (GREEN_FILL if is_best else NO_FILL)

    for c, val in enumerate([
        f"{res['dx_fixed']} u",
        sim["setup_cost"], sim["holding_cost"], sim["backorder_cost"],
        ot_x_total, ot_y_total, res["invest_X"], res["invest_Y"],
        sim["total_cost"],
    ], start=2):
        is_money = isinstance(val, float)
        plain(ws.cell(r, c), val,
              bold=(c == 10 and is_best),
              align="right" if is_money else "center",
              fmt='"€"#,##0.00' if is_money else None,
              border=full_thin, fill=row_fill,
              color="CC0000" if (c == 5 and is_money and val > 0) else "000000")

#  Backorders per period 
r += 2
BO_LAST_COL = 2 + T + 1        
section_title(ws, r, BO_LAST_COL, "Backorders per period (units, realized demand)")
r += 1
ws.row_dimensions[r].height = 18
hdr(ws.cell(r, 2), "dx (units)")
for t in periods:
    col = 2 + t
    ws.column_dimensions[get_column_letter(col)].width = 8
    hdr(ws.cell(r, col), f"P{t}")
hdr(ws.cell(r, BO_LAST_COL), "Total units")

for res in results:
    r += 1
    ws.row_dimensions[r].height = 16
    sim     = res["sim"]
    is_ref  = (res["dx_fixed"] == REF_DX)
    is_best = abs(sim["total_cost"] - best_realized) < 0.01
    row_fill = LGREY_FILL if is_ref else (GREEN_FILL if is_best else NO_FILL)

    plain(ws.cell(r, 2), f"{res['dx_fixed']} u",
          align="left", border=full_thin, fill=row_fill)
    total_bo = 0.0
    for t in periods:
        col = 2 + t
        bo  = sim["backorders"][t]
        total_bo += bo
        plain(ws.cell(r, col),
              round(bo, 2) if bo > 0 else "",
              align="right",
              fmt="#,##0.00" if bo > 0 else None,
              border=full_thin, fill=row_fill,
              color="CC0000" if bo > 0 else "000000")
    plain(ws.cell(r, BO_LAST_COL),
          round(total_bo, 2) if total_bo > 0 else "",
          bold=True, align="right",
          fmt="#,##0.00" if total_bo > 0 else None,
          border=full_thin, fill=row_fill,
          color="CC0000" if total_bo > 0 else "000000")
ws.column_dimensions[get_column_letter(BO_LAST_COL)].width = 12

#  Chart data + charts 
r += 2
section_title(ws, r, LAST_COL, "Charts")
cds = r + 1

for c, h in enumerate(["dx (units)", "Forecast total",
                        "Realized total", "Backorder cost", "Holding cost"], start=2):
    cell = ws.cell(cds, c)
    cell.value = h
    cell.font  = Font(name="Calibri", bold=True, size=9, color="FFFFFF")
    cell.fill  = BLACK_FILL
    cell.alignment = Alignment(horizontal="center")

for idx, res in enumerate(results):
    row   = cds + 1 + idx
    sim   = res["sim"]
    rfill = LGREY_FILL if res["dx_fixed"] == REF_DX else NO_FILL
    for c, val in enumerate([
        f"{res['dx_fixed']} u",
        round(res["forecast_total"], 2),
        round(sim["total_cost"], 2),
        round(sim["backorder_cost"], 2),
        round(sim["holding_cost"], 2),
    ], start=2):
        cell = ws.cell(row, c)
        cell.value     = val
        cell.font      = Font(name="Calibri", size=9)
        cell.alignment = Alignment(horizontal="center")
        cell.border    = full_thin
        cell.fill      = rfill

n   = len(results)
cde = cds + n
lbl = Reference(ws, min_col=2, max_col=2, min_row=cds+1, max_row=cde)

# Chart 1: forecast vs realized total cost
chart1 = BarChart()
chart1.type = "col"; chart1.grouping = "clustered"
chart1.title = "Total Cost per X-Expansion Scenario"
chart1.y_axis.title = "Total Cost (€)"
chart1.x_axis.title = "Permanent expansion dx (units)"
chart1.style = 10; chart1.width = 24; chart1.height = 14
s1 = Series(Reference(ws, min_col=3, max_col=3, min_row=cds, max_row=cde),
            title="Forecast total")
s1.graphicalProperties.solidFill = "4472C4"
s2 = Series(Reference(ws, min_col=4, max_col=4, min_row=cds, max_row=cde),
            title="Realized total")
s2.graphicalProperties.solidFill = "ED7D31"
chart1.append(s1); chart1.append(s2); chart1.set_categories(lbl)
chart1.x_axis.tickLblPos = "low"
chart1.x_axis.noMultiLvlLbl = True
chart1.x_axis.numFmt = "General"
chart1.x_axis.tickMarkSkip = 1
ws.add_chart(chart1, f"B{cde + 2}")

# Chart 2: holding vs backorder (stacked)
chart2 = BarChart()
chart2.type = "col"; chart2.grouping = "stacked"
chart2.title = "Holding vs Backorder Cost under Realized Demand"
chart2.y_axis.title = "Cost (€)"
chart2.x_axis.title = "Permanent expansion dx (units)"
chart2.style = 10; chart2.width = 24; chart2.height = 14
s3 = Series(Reference(ws, min_col=6, max_col=6, min_row=cds, max_row=cde),
            title="Holding cost")
s3.graphicalProperties.solidFill = "4472C4"
s4 = Series(Reference(ws, min_col=5, max_col=5, min_row=cds, max_row=cde),
            title="Backorder cost")
s4.graphicalProperties.solidFill = "FF0000"
chart2.append(s3); chart2.append(s4); chart2.set_categories(lbl)
chart2.x_axis.tickLblPos = "low"
chart2.x_axis.noMultiLvlLbl = True
chart2.x_axis.numFmt = "General"
chart2.x_axis.tickMarkSkip = 1
ws.add_chart(chart2, f"B{cde + 24}")


wb.save(OUTPUT_FILE)
print(f"\nResults written to {OUTPUT_FILE} -> sheet '{SHEET_NAME}'")